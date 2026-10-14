import copy,hashlib,json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
from scripts import normalize_public_content as n
from scripts.msds_register_qc import check_product_boundaries,history_identity_issues
from scripts.validate_public_release import ValidationResult,validate_cas_fields
from scripts.cas_numbers import SOURCE_IDENTIFIER
from scripts.convert_excel_to_json import make_ingredient

ROOT=Path(__file__).resolve().parents[1]
class CompanyMsdsRepairTests(unittest.TestCase):
    def test_source_values_preserved_and_not_grouped_as_cas(self):
        ps=json.loads((ROOT/'data/msds.public.json').read_text(encoding='utf-8'))
        repairs=json.loads((ROOT/'data/msds-content-repairs.json').read_text(encoding='utf-8'))
        for fix in repairs['ingredientIdentifiers']:
            if fix['beforeCas']=='2028-08-3':continue
            p=next(p for p in ps if p['id']==fix['productId'])
            row=next(r for r in p['ingredients'] if r['chemicalName']==fix['chemicalName'])
            self.assertEqual(row['casNo'],'')
            self.assertEqual(row[fix['sourceField']],fix['beforeCas'])
            self.assertEqual(hashlib.sha256((ROOT/fix['pdf']).read_bytes()).hexdigest(),fix['sha256'])
    def test_changed_pdf_does_not_receive_old_identifier_fix(self):
        with tempfile.TemporaryDirectory()as d:
            p=Path(d);(p/'x.pdf').write_bytes(b'new')
            products=[{'id':'x','ingredients':[{'chemicalName':'a','casNo':'99-3-1171','content':'1'}]}]
            fix={'productId':'x','chemicalName':'a','beforeCas':'99-3-1171','sourceIdentifier':'99-3-1171','pdf':'x.pdf','sha256':'old'}
            with patch.object(n,'ROOT',p):n.apply_repairs(products,[],{'ingredientIdentifiers':[fix]}, {})
            self.assertEqual(products[0]['ingredients'][0]['casNo'],'99-3-1171')
    def test_identifier_fix_does_not_change_other_chemical(self):
        products=[{'id':'x','ingredients':[{'chemicalName':'b','casNo':'99-3-1171'}]}]
        fix={'productId':'x','chemicalName':'a','beforeCas':'99-3-1171','sourceIdentifier':'99-3-1171'}
        n.apply_repairs(products,[],{'ingredientIdentifiers':[fix]}, {})
        self.assertEqual(products[0]['ingredients'][0]['casNo'],'99-3-1171')
    def test_identifier_fix_does_not_overwrite_another_reviewed_identifier(self):
        products=[{'id':'x','ingredients':[{'chemicalName':'a','casNo':'','sourceIdentifier':'KE-222'}]}]
        fix={'productId':'x','chemicalName':'a','beforeCas':'99-3-1171','sourceIdentifier':'99-3-1171'}
        n.apply_repairs(products,[],{'ingredientIdentifiers':[fix]}, {})
        self.assertEqual(products[0]['ingredients'][0]['sourceIdentifier'],'KE-222')
    def test_candidate_has_six_source_rows(self):
        p=next(x for x in json.loads((ROOT/'data/msds.public.json').read_text(encoding='utf-8')) if x['id']=='msds-pdf-f60438c81a0521d8')
        o=next(x for x in json.loads((ROOT/'data/msds-overrides.public.json').read_text(encoding='utf-8')) if x.get('sourcePdfPath')==p['pdfPath'])
        self.assertEqual(o['ingredients'],p['ingredients']);self.assertEqual(len(o['ingredients']),6)
        self.assertEqual(o['ingredients'][-1]['casNo'],'128-37-0')
    def test_source_error_is_still_reported_and_cannot_have_active_cas(self):
        result=ValidationResult();validate_cas_fields(result,{'ingredients':[{'casNo':'','sourceCasNo':'9000-11-07'}]},'x')
        self.assertEqual(result.warnings[0].code,'SOURCE_CAS_ERROR')
        result=ValidationResult();validate_cas_fields(result,{'ingredients':[{'casNo':'9004-32-4','sourceCasNo':'9000-11-07'}]},'x')
        self.assertIn('SOURCE_CAS_ERROR_IN_ACTIVE_CAS',[e.code for e in result.errors])
    def test_supplier_scope_and_identity_pending(self):
        r=json.loads((ROOT/'data/msds-register.json').read_text(encoding='utf-8'))['products']['msds-pdf-d462ab4576452226']
        self.assertEqual(r['supplier'],'한일루켐(주)');self.assertEqual(r['status'],'pending')
        self.assertEqual(r['supplierEvidence']['scope'],'linked_pdf');self.assertEqual(r['supplierEvidence']['productIdentityStatus'],'pending')
        self.assertIn('EP 00',r['sourceIssue'])
    def test_qc_detects_collapsed_blocks_without_labels(self):
        w=Workbook();s=w.active;s['E4']='first';s.merge_cells('E4:E7')
        evidence={'blocks':[{'start':4,'end':5},{'start':6,'end':7}]}
        issues=check_product_boundaries(s,evidence)
        self.assertIn(('PRODUCT_BOUNDARY_LOST',6,7,'E'),issues)
        self.assertIn(('BOUNDARY_IDENTITY_MISSING',6,7),issues)
    def test_qc_checks_composition_against_boundary_evidence(self):
        w=Workbook();s=w.active;s['T4']='67-64-1';s['U4']='5'
        issues=check_product_boundaries(s,{'blocks':[{'start':4,'end':4,'casContent':[['67-64-1','10']]}]})
        self.assertIn(('BOUNDARY_COMPONENT_MISMATCH',4,4),issues)
    def test_qc_rejects_boundary_evidence_for_replaced_pdf(self):
        with tempfile.TemporaryDirectory()as d:
            root=Path(d);(root/'source.pdf').write_bytes(b'new')
            w=Workbook();s=w.active
            issues=check_product_boundaries(s,{'blocks':[{'start':4,'end':4,'pdf':'source.pdf','sha256':'old'}]},root)
            self.assertIn(('BOUNDARY_SOURCE_STALE',4,4,'source.pdf'),issues)
    def test_qc_reports_new_product_lost_from_history(self):
        w=Workbook();s=w.active;s['H4']='first.pdf';h=w.create_sheet('업데이트_이력')
        h.append(['file','name','no','action']);h.append(['missing.pdf','lost','','신규추가(계속누적)'])
        self.assertEqual(history_identity_issues(w,s),[('HISTORY_PRODUCT_IDENTITY_MISSING',2,'missing.pdf')])
    def test_qc_accepts_original_register_filename_only_with_verified_rename(self):
        with tempfile.TemporaryDirectory()as d:
            root=Path(d);(root/'data').mkdir();(root/'new.pdf').write_bytes(b'original source')
            (root/'data/pdf-renames.json').write_text(json.dumps({'renames':{'old.pdf':'new.pdf'}}),encoding='utf-8')
            w=Workbook();s=w.active;s['E4']='product';s['H4']='old.pdf'
            evidence={'blocks':[{'start':4,'end':4,'pdf':'new.pdf','registerPdf':'old.pdf','sha256':hashlib.sha256(b'original source').hexdigest()}]}
            self.assertNotIn(('BOUNDARY_SOURCE_FILE_MISMATCH',4,4),check_product_boundaries(s,evidence,root))
            (root/'data/pdf-renames.json').write_text(json.dumps({'renames':{'old.pdf':'another.pdf'}}),encoding='utf-8')
            self.assertIn(('BOUNDARY_SOURCE_FILE_MISMATCH',4,4),check_product_boundaries(s,evidence,root))
    def test_qc_does_not_accept_rename_when_source_content_changed(self):
        with tempfile.TemporaryDirectory()as d:
            root=Path(d);(root/'data').mkdir();(root/'new.pdf').write_bytes(b'replaced source')
            (root/'data/pdf-renames.json').write_text(json.dumps({'renames':{'old.pdf':'new.pdf'}}),encoding='utf-8')
            w=Workbook();s=w.active;s['E4']='product';s['H4']='old.pdf'
            evidence={'blocks':[{'start':4,'end':4,'pdf':'new.pdf','registerPdf':'old.pdf','sha256':hashlib.sha256(b'original source').hexdigest()}]}
            issues=check_product_boundaries(s,evidence,root)
            self.assertIn(('BOUNDARY_SOURCE_STALE',4,4,'new.pdf'),issues)
            self.assertIn(('BOUNDARY_SOURCE_FILE_MISMATCH',4,4),issues)
    def test_excel_conversion_keeps_distinct_source_fields(self):
        row=make_ingredient({'chemicalName':'a','content':'1','casNo':'CAS 미기재','sourceIdentifier':'701-048-1'})
        self.assertEqual(row['sourceIdentifier'],'701-048-1');self.assertEqual(row['casNo'],'')
        row=make_ingredient({'chemicalName':'a','content':'1','casNo':'원문 CAS 오류','sourceIdentifier':'9000-11-07','casSourceIssue':'원문 CAS번호 오류·정정본 확인'})
        self.assertEqual(row['sourceCasNo'],'9000-11-07');self.assertNotIn('sourceIdentifier',row)
        row=make_ingredient({'chemicalName':'a','content':'1','casNo':'CAS 확인 필요','sourceIdentifier':'KE-08612','casSourceIssue':'기존 등록대장 식별값 보존·연결 원문 대응 미확인'})
        self.assertEqual(row['sourceIdentifierStatus'],'unverified_register')
    def test_parser_recognizes_opaque_identifiers_without_cas_truncation(self):
        for value in ['2015-3-7140','99-3-1171','701-048-1','KE-28667']:
            self.assertEqual(SOURCE_IDENTIFIER.search('a '+value+' 1~2').group(),value)
        self.assertIsNone(SOURCE_IDENTIFIER.search('9000-11-07'))
    def test_restored_boundary_products_keep_correct_composition_ranges(self):
        products=json.loads((ROOT/'data/msds.public.json').read_text(encoding='utf-8'))
        hr=next(p for p in products if p['id']=='msds-pdf-3b92d1fa1c9a3177')
        self.assertEqual([(r['casNo'],r['content'])for r in hr['ingredients']],[('9003-07-0','8~12'),('9010-79-1','50~60'),('25087-34-7','10~15'),('14807-96-6','15~20'),('31570-04-4','0.1~1'),('167078-06-0','0.1~1')])
        therm=next(p for p in products if p['id']=='msds-pdf-48f0afe382fbccab')
        self.assertEqual(len(therm['ingredients']),2)
        self.assertEqual(therm['ingredients'][0]['content'],'98~99')
        self.assertEqual(therm['ingredients'][0]['sourceIdentifier'],'KE-12546')
        self.assertEqual(therm['ingredients'][1]['content'],'0.1~1')
