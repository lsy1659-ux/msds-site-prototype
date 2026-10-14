from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, PatternFill
from openpyxl.utils import get_column_letter


PRODUCT_COLS = list(range(1, 19)) + list(range(29, 49))  # A:R, AC:AV
COMPONENT_COLS = list(range(19, 29)) + list(range(49, 51))  # S:AB, AW:AX
TARGET_COLS = list(range(22, 29))  # V:AB
GROUP_START_COLS = [1, 2, 5, 8]  # index, MSDS no., product name, file name
CHECK_COLS = list(range(1, 51))

GREEN = "FFC6EFCE"
PLACEHOLDERS = {"name", "cas", "pct"}
ERROR_VALUES = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}
CIRCLE_VALUES = {"○", "O", "o", "0"}


@dataclass
class ProductGroup:
    start: int
    end: int
    label: str

    @property
    def size(self) -> int:
        return self.end - self.start + 1


def normalize_text(value) -> str:
    return "" if value is None else str(value).strip()


def rgb(cell) -> str:
    fill = cell.fill
    if fill and fill.fill_type == "solid" and fill.fgColor and fill.fgColor.type == "rgb":
        return (fill.fgColor.rgb or "").upper()
    return ""


def is_green(cell) -> bool:
    return rgb(cell) == GREEN


def row_has_value(ws, row: int, cols: Iterable[int]) -> bool:
    return any(normalize_text(ws.cell(row, col).value) for col in cols)


def find_main_sheet(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10)):
            values = [normalize_text(cell.value) for cell in row]
            joined = " ".join(values)
            if "화학물질명" in joined and "CAS" in joined and "함유" in joined:
                return ws
    return wb.active


def find_header_row(ws) -> int:
    for row in range(1, min(ws.max_row, 10) + 1):
        values = " ".join(normalize_text(ws.cell(row, col).value) for col in range(1, min(ws.max_column, 55) + 1))
        if "화학물질명" in values and "CAS" in values and "함유" in values:
            return row
    return 3


def find_last_data_row(ws, data_start: int) -> int:
    for row in range(ws.max_row, data_start - 1, -1):
        if row_has_value(ws, row, CHECK_COLS):
            return row
    return data_start - 1


def is_group_start(ws, row: int) -> bool:
    values = [normalize_text(ws.cell(row, col).value) for col in GROUP_START_COLS]
    if not any(values):
        return False
    # Avoid treating extracted component-only rows as product starts.
    if values[1].lower() in PLACEHOLDERS or values[3].lower() in PLACEHOLDERS:
        return False
    return True


def product_label(ws, row: int) -> str:
    for col in [2, 5, 8]:
        value = normalize_text(ws.cell(row, col).value)
        if value:
            return value
    return f"row {row}"


def find_product_groups(ws) -> list[ProductGroup]:
    data_start = find_header_row(ws) + 1
    last_row = find_last_data_row(ws, data_start)
    starts = [row for row in range(data_start, last_row + 1) if is_group_start(ws, row)]
    groups: list[ProductGroup] = []
    for idx, start in enumerate(starts):
        stop = starts[idx + 1] - 1 if idx + 1 < len(starts) else last_row
        while stop > start and not row_has_value(ws, stop, CHECK_COLS):
            stop -= 1
        groups.append(ProductGroup(start, stop, product_label(ws, start)))
    return groups


def merged_full(ws, start: int, end: int, col: int) -> bool:
    if start == end:
        return True
    for merged_range in ws.merged_cells.ranges:
        if (
            merged_range.min_col == col
            and merged_range.max_col == col
            and merged_range.min_row == start
            and merged_range.max_row == end
        ):
            return True
    return False


def unmerge_overlaps(ws, start: int, end: int, cols: Iterable[int]) -> None:
    col_set = set(cols)
    to_unmerge = []
    for merged_range in list(ws.merged_cells.ranges):
        range_cols = set(range(merged_range.min_col, merged_range.max_col + 1))
        if range_cols & col_set and not (merged_range.max_row < start or merged_range.min_row > end):
            to_unmerge.append(str(merged_range))
    for ref in to_unmerge:
        ws.unmerge_cells(ref)


def copy_style(source, target) -> None:
    if source.has_style:
        target._style = copy.copy(source._style)
    target.number_format = source.number_format
    target.font = copy.copy(source.font)
    target.fill = copy.copy(source.fill)
    target.border = copy.copy(source.border)
    target.alignment = copy.copy(source.alignment)
    target.protection = copy.copy(source.protection)


def normalize_product_merges(ws, groups: list[ProductGroup]) -> int:
    changed = 0
    for group in groups:
        if group.size <= 1:
            continue
        unmerge_overlaps(ws, group.start, group.end, PRODUCT_COLS)
        for col in PRODUCT_COLS:
            top = ws.cell(group.start, col)
            top.alignment = Alignment(
                horizontal=top.alignment.horizontal or "center",
                vertical="center",
                wrap_text=True,
            )
            for row in range(group.start + 1, group.end + 1):
                cell = ws.cell(row, col)
                copy_style(top, cell)
                cell.value = None
            ws.merge_cells(start_row=group.start, start_column=col, end_row=group.end, end_column=col)
            changed += 1
    return changed


def normalize_target_flags(ws, data_start: int, last_row: int) -> int:
    changed = 0
    green_fill = PatternFill(fill_type="solid", fgColor=GREEN)
    blank_fill = PatternFill(fill_type=None)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in range(data_start, last_row + 1):
        for col in TARGET_COLS:
            cell = ws.cell(row, col)
            text = normalize_text(cell.value)
            if text in CIRCLE_VALUES:
                if cell.value != "○" or not is_green(cell):
                    changed += 1
                cell.value = "○"
                cell.fill = green_fill
                cell.alignment = center
            elif text == "":
                if is_green(cell):
                    changed += 1
                cell.fill = blank_fill
                cell.alignment = center
    return changed


def check_product_boundaries(ws, evidence: dict, source_root: Path | None = None) -> list:
    """Validate independent source-confirmed blocks, even when start labels were lost."""
    issues = []
    for block in evidence.get("blocks", []):
        start, end = int(block["start"]), int(block["end"])
        source_verified = False
        if source_root is not None and block.get("pdf") and block.get("sha256"):
            source=source_root / block["pdf"]
            source_verified = source.is_file() and hashlib.sha256(source.read_bytes()).hexdigest()==block["sha256"]
            if not source_verified:
                issues.append(("BOUNDARY_SOURCE_STALE",start,end,block["pdf"]))
        for col in PRODUCT_COLS:
            if not merged_full(ws, start, end, col):
                issues.append(("PRODUCT_BOUNDARY_LOST", start, end, get_column_letter(col)))
        expected = block.get("casContent") or []
        actual = [[normalize_text(ws.cell(row,20).value), normalize_text(ws.cell(row,21).value)] for row in range(start,end+1)]
        if expected and actual != expected:
            issues.append(("BOUNDARY_COMPONENT_MISMATCH", start, end))
        if not normalize_text(ws.cell(start,5).value) or not normalize_text(ws.cell(start,8).value):
            issues.append(("BOUNDARY_IDENTITY_MISSING", start, end))
        if block.get("pdf"):
            allowed_names = {Path(block["pdf"]).name}
            original_pdf = block.get("registerPdf")
            if source_verified and original_pdf:
                rename_file = source_root / "data" / "pdf-renames.json"
                if rename_file.is_file():
                    renames = json.loads(rename_file.read_text(encoding="utf-8")).get("renames") or {}
                    if renames.get(original_pdf) == block["pdf"]:
                        allowed_names.add(Path(original_pdf).name)
            if normalize_text(ws.cell(start,8).value) not in allowed_names:
                issues.append(("BOUNDARY_SOURCE_FILE_MISMATCH",start,end))
    return issues


def history_identity_issues(wb, ws) -> list:
    """New-product histories with lost filename cells must not disappear in a green QC."""
    if "업데이트_이력" not in wb.sheetnames:
        return []
    files = {normalize_text(ws.cell(row,8).value) for row in range(4,ws.max_row+1)} - {""}
    issues=[]
    for row in wb["업데이트_이력"].iter_rows(min_row=2):
        if len(row)<4 or normalize_text(row[3].value)!="신규추가(계속누적)":
            continue
        name=normalize_text(row[0].value)
        if name and name not in files:
            issues.append(("HISTORY_PRODUCT_IDENTITY_MISSING",row[0].row,name))
    return issues


def scan_workbook(path: Path, fix: bool, output: Path | None, boundary_evidence: dict | None = None, source_root: Path | None = None) -> dict:
    wb = load_workbook(path)
    ws = find_main_sheet(wb)
    header_row = find_header_row(ws)
    data_start = header_row + 1
    last_row = find_last_data_row(ws, data_start)
    groups = find_product_groups(ws)

    placeholder_hits = []
    error_hits = []
    merge_issues = []
    flag_issues = []
    boundary_issues = check_product_boundaries(ws,boundary_evidence or {},source_root) + history_identity_issues(wb,ws)
    identifier_issues=[]

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                text = normalize_text(cell.value)
                if cell.data_type == "e" or text in ERROR_VALUES:
                    error_hits.append((sheet.title, cell.coordinate, text))

    for row in range(data_start, last_row + 1):
        for col in [19, 20, 21]:
            text = normalize_text(ws.cell(row, col).value).lower()
            if text in PLACEHOLDERS:
                placeholder_hits.append((row, get_column_letter(col), text))
        cas=normalize_text(ws.cell(row,20).value)
        if re.match(r"(?:KE-|T-|Z\d+-)",cas):
            identifier_issues.append(("SOURCE_IDENTIFIER_IN_CAS_COLUMN",row,cas))
        if re.fullmatch(r"\d+(?:-\d+){2}",cas):
            digits=cas.replace("-","")
            standard=bool(re.fullmatch(r"\d{2,7}-\d{2}-\d",cas))
            valid=standard and sum(int(d)*i for i,d in enumerate(reversed(digits[:-1]),1))%10==int(digits[-1])
            if not valid:identifier_issues.append(("NON_CAS_IN_CAS_COLUMN",row,cas))

    for group in groups:
        if group.size <= 1:
            continue
        for col in PRODUCT_COLS:
            if not merged_full(ws, group.start, group.end, col):
                merge_issues.append((group.label, group.start, group.end, get_column_letter(col)))

    for row in range(data_start, last_row + 1):
        for col in TARGET_COLS:
            cell = ws.cell(row, col)
            text = normalize_text(cell.value)
            if text in CIRCLE_VALUES and not is_green(cell):
                flag_issues.append((row, get_column_letter(col), "circle_without_green"))
            if text == "" and is_green(cell):
                flag_issues.append((row, get_column_letter(col), "green_without_circle"))

    fixed_merges = 0
    fixed_flags = 0
    if fix:
        if boundary_issues:
            raise ValueError("원문 제품 경계·이력 불일치가 있어 자동 병합을 중단합니다. 확인한 제품별 경계와 값을 먼저 복구하세요.")
        fixed_merges = normalize_product_merges(ws, groups)
        fixed_flags = normalize_target_flags(ws, data_start, last_row)
        save_path = output or path.with_name(path.stem + "_qc_fixed.xlsx")
        wb.save(save_path)
    else:
        save_path = None

    return {
        "sheet": ws.title,
        "groups": len(groups),
        "multi_row_groups": sum(1 for group in groups if group.size > 1),
        "placeholder_hits": placeholder_hits,
        "error_hits": error_hits,
        "merge_issues": merge_issues,
        "flag_issues": flag_issues,
        "boundary_issues": boundary_issues,
        "identifier_issues": identifier_issues,
        "source_review_rows": [row for row in range(data_start,last_row+1) if any(t in normalize_text(ws.cell(row,52).value) for t in ["미확인","정정본"])],
        "fixed_merges": fixed_merges,
        "fixed_flags": fixed_flags,
        "saved": str(save_path) if save_path else "",
    }


def print_report(report: dict) -> None:
    print(f"SHEET={report['sheet']}")
    print(f"GROUPS={report['groups']}")
    print(f"MULTI_ROW_GROUPS={report['multi_row_groups']}")
    print(f"PLACEHOLDER_HITS={len(report['placeholder_hits'])}")
    print(f"ERROR_HITS={len(report['error_hits'])}")
    print(f"MERGE_ISSUES={len(report['merge_issues'])}")
    print(f"FLAG_ISSUES={len(report['flag_issues'])}")
    print(f"BOUNDARY_ISSUES={len(report['boundary_issues'])}")
    print(f"IDENTIFIER_ISSUES={len(report['identifier_issues'])}")
    print(f"SOURCE_REVIEW_ROWS={report['source_review_rows']}")
    print(f"FIXED_MERGES={report['fixed_merges']}")
    print(f"FIXED_FLAGS={report['fixed_flags']}")
    if report["saved"]:
        print(f"SAVED={report['saved']}")
    for key in ["placeholder_hits", "error_hits", "merge_issues", "flag_issues", "boundary_issues", "identifier_issues"]:
        for item in report[key][:25]:
            print(f"{key.upper()} {item}")


def main() -> None:
    parser = argparse.ArgumentParser(description="QC and normalize CAMS MSDS register workbooks.")
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--fix", action="store_true", help="normalize product merges and V:AB target flag fills")
    parser.add_argument("--output", type=Path, help="output workbook path when --fix is used")
    parser.add_argument("--boundary-evidence", type=Path, help="원문 대조로 확인한 제품별 행 경계와 CAS/함유량 JSON")
    args = parser.parse_args()

    evidence=json.loads(args.boundary_evidence.read_text(encoding="utf-8")) if args.boundary_evidence else None
    source_root=args.boundary_evidence.resolve().parents[1] if args.boundary_evidence else None
    report = scan_workbook(args.workbook, args.fix, args.output,evidence,source_root)
    print_report(report)
    if any(report[k] for k in ["placeholder_hits","error_hits","merge_issues","flag_issues","boundary_issues","identifier_issues"]):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
